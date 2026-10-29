# Convierte "Conglomerado RL.xlsx" en datos/conglomerado-rl.js, que todas las paginas cargan por defecto.
# Uso (desde la carpeta del proyecto):  python herramientas/excel_a_js.py ["ruta/al/archivo.xlsx"]
# Requiere openpyxl (pip install openpyxl). Se conservan TODAS las columnas y valores de la primera hoja.

import datetime
import json
import os
import shutil
import sys
import tempfile

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ENTRADA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(RAIZ, 'Conglomerado RL.xlsx')
SALIDA = os.path.join(RAIZ, 'datos', 'conglomerado-rl.js')


def texto(v):
    # Mismo criterio que la lectura anterior en el navegador: fechas AAAA-MM-DD, enteros sin decimales
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        f = v.strftime('%Y-%m-%d')
        return f + v.strftime(' %H:%M') if (v.hour or v.minute) else f
    if isinstance(v, datetime.date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def main():
    # se trabaja sobre una copia: Excel u OneDrive pueden tener bloqueado el original
    tmp = os.path.join(tempfile.mkdtemp(), 'conglomerado.xlsx')
    shutil.copyfile(ENTRADA, tmp)
    libro = openpyxl.load_workbook(tmp, read_only=True, data_only=True)
    nombre_hoja = libro.sheetnames[0]
    filas = [[texto(c) for c in r] for r in libro[nombre_hoja].iter_rows(values_only=True)]
    libro.close()
    columnas = [c.strip() for c in filas[0]]
    datos = [r + [''] * (len(columnas) - len(r)) for r in filas[1:] if any(x != '' for x in r)]
    meta = {
        'archivo': os.path.basename(ENTRADA), 'hojaExcel': nombre_hoja,
        'cargado': datetime.datetime.now().isoformat(timespec='seconds'), 'filas': len(datos)
    }
    d = lambda x: json.dumps(x, ensure_ascii=False, separators=(',', ':'))
    with open(SALIDA, 'w', encoding='utf-8', newline='\n') as f:
        f.write('// Generado por herramientas/excel_a_js.py a partir de ' + meta['archivo'] + '. No editar a mano.\n')
        f.write('var CONGLOMERADO_RL = {\n"meta":' + d(meta) + ',\n"columnas":' + d(columnas) + ',\n"filas":[\n')
        f.write(',\n'.join(d(r) for r in datos))
        f.write('\n]};\n')
    print('%s: %d filas, %d columnas -> %s' % (meta['archivo'], len(datos), len(columnas), SALIDA))


if __name__ == '__main__':
    main()
